"""Valida que un Excel de cartera respete la estructura fija de columnas.

Uso:  python3 validar_cartera.py ruta/a/cartera.xlsx [nombre_hoja]

Compara la fila 1 con columnas-cartera.json (mismos nombres, mismo orden)
y revisa que cada fila con datos tenga Customer ID. No modifica el archivo.
Sale con código 0 si la estructura es correcta y 1 si no.
"""
import json
import sys
import unicodedata
from pathlib import Path

import openpyxl

ESQUEMA = Path(__file__).with_name("columnas-cartera.json")


def norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(s.lower().split())


def validar(ruta, hoja=None):
    esquema = json.loads(ESQUEMA.read_text(encoding="utf-8"))
    esperadas = [c["nombre"] for c in esquema["columnas"]]
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[hoja] if hoja else (wb[esquema["hoja"]] if esquema["hoja"] in wb.sheetnames else wb.worksheets[0])

    filas = ws.iter_rows(values_only=True)
    encabezado = [str(v).strip() if v is not None else "" for v in next(filas, ())]
    while encabezado and not encabezado[-1]:
        encabezado.pop()

    errores, avisos = [], []
    reales = [norm(h) for h in encabezado]
    for i, nombre in enumerate(esperadas):
        if norm(nombre) not in reales:
            errores.append(f"Falta la columna «{nombre}».")
        elif reales.index(norm(nombre)) != i:
            errores.append(f"«{nombre}» está en la columna {reales.index(norm(nombre)) + 1}; debe estar en la {i + 1}.")
        elif encabezado[i] != nombre:
            avisos.append(f"«{encabezado[i]}» se escribe distinto a «{nombre}» (acentos o mayúsculas).")
    for h in encabezado:
        if h and norm(h) not in [norm(e) for e in esperadas]:
            errores.append(f"Columna no prevista: «{h}».")

    clientes = sin_id = 0
    if not errores:
        i_id = esperadas.index(esquema["clave_cliente"])
        for fila in filas:
            if not any(v not in (None, "") for v in fila):
                continue
            clientes += 1
            if i_id >= len(fila) or fila[i_id] in (None, ""):
                sin_id += 1
        if sin_id:
            avisos.append(f"{sin_id} fila(s) con datos no tienen Customer ID.")

    print(f"Hoja: {ws.title}")
    print(f"Columnas: {len(encabezado)} de {len(esperadas)} esperadas")
    print(f"Clientes (filas con datos): {clientes}")
    for e in errores:
        print("ERROR  " + e)
    for a in avisos:
        print("AVISO  " + a)
    print("Estructura correcta." if not errores else "La estructura no coincide; no se debe procesar este archivo.")
    return not errores


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    sys.exit(0 if validar(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None) else 1)
